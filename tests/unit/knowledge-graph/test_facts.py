#!/usr/bin/env python3
"""Tests for Sheet 7 (Facts) parsing. Spec §4.5."""
import importlib.util, os, pytest
from openpyxl import Workbook


def load_parser():
    script_path = os.path.join(os.path.dirname(__file__), "../../../scripts/generate-knowledge-graph.py")
    spec = importlib.util.spec_from_file_location("parser", os.path.abspath(script_path))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

parser = load_parser()


def make_facts_sheet(wb, rows):
    ws = wb.create_sheet("Facts")
    ws.append(["Facts"])
    ws.append(["Clinical fact vocabulary"])
    ws.append([])
    ws.append(["Machine Name", "Display Name", "Type", "Unit",
               "Valid For", "Required", "Area",
               "Applicable Classifications", "Applicable Phases",
               "Extraction Hint"])
    for r in rows:
        ws.append(r)
    return ws


# Valid base row — all fields correct
BASE_ROW = ["weight", "Weight", "number", "kg", "7d", "Yes", "vitals", "ALL", "ALL", "Patient weight"]


class TestFacts:

    def test_valid_sheet_yields_dict_keyed_by_machine_name(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_facts_sheet(wb, [BASE_ROW])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert errors == []
        assert "weight" in facts
        f = facts["weight"]
        assert set(f.keys()) >= {"display_name", "type", "unit", "valid_for", "valid_for_hours",
                                  "required", "area", "applicable_classifications",
                                  "applicable_phases", "extraction_hint"}
        assert f["type"] == "number"
        assert f["unit"] == "kg"
        assert f["required"] == True
        assert f["valid_for"] == "7d"
        assert f["valid_for_hours"] == 168

    def test_absent_sheet_returns_empty_dict(self):
        facts, errors = parser.parse_facts(None, ["CLS1"], ["phase1"])
        assert facts == {}
        assert errors == []

    def test_duplicate_machine_name_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_facts_sheet(wb, [BASE_ROW, BASE_ROW])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert len(errors) >= 1
        assert any("weight" in e["error"] for e in errors)

    def test_number_type_without_unit_is_rejected(self):
        row = list(BASE_ROW)
        row[3] = ""  # Unit column (index 3) empty
        wb = Workbook(); wb.remove(wb.active)
        ws = make_facts_sheet(wb, [row])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert len(errors) >= 1
        assert any("Unit" in e["error"] for e in errors)

    def test_invalid_type_is_rejected(self):
        row = list(BASE_ROW)
        row[2] = "invalid_type"  # Type column (index 2)
        wb = Workbook(); wb.remove(wb.active)
        ws = make_facts_sheet(wb, [row])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert len(errors) >= 1
        assert any("invalid_type" in e["error"] for e in errors)

    def test_invalid_area_is_rejected(self):
        row = list(BASE_ROW)
        row[6] = "invalid_area"  # Area column (index 6)
        wb = Workbook(); wb.remove(wb.active)
        ws = make_facts_sheet(wb, [row])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert len(errors) >= 1
        assert any("invalid_area" in e["error"] or "Area" in e.get("column", "") for e in errors)

    def test_malformed_valid_for_is_rejected(self):
        row = list(BASE_ROW)
        row[4] = "3 days"  # Valid For column (index 4) — space makes it invalid
        wb = Workbook(); wb.remove(wb.active)
        ws = make_facts_sheet(wb, [row])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert len(errors) >= 1
        assert any("Valid For" in e["error"] or "3 days" in e["error"] for e in errors)

    def test_unknown_classification_is_rejected(self):
        row = list(BASE_ROW)
        row[7] = "UNKNOWN_CLS"  # Applicable Classifications column (index 7)
        wb = Workbook(); wb.remove(wb.active)
        ws = make_facts_sheet(wb, [row])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert len(errors) >= 1
        assert any("UNKNOWN_CLS" in e["error"] for e in errors)

    def test_unknown_phase_is_rejected(self):
        row = list(BASE_ROW)
        row[8] = "UNKNOWN_PHASE"  # Applicable Phases column (index 8)
        wb = Workbook(); wb.remove(wb.active)
        ws = make_facts_sheet(wb, [row])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert len(errors) >= 1
        assert any("UNKNOWN_PHASE" in e["error"] for e in errors)

    def test_bad_classification_excludes_row_from_facts(self):
        """A row with an unknown classification must not appear in the output dict."""
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW)
        row[7] = "UNKNOWN_CLS"  # Applicable Classifications
        ws = make_facts_sheet(wb, [row])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert len(errors) >= 1
        assert "weight" not in facts  # the row must be excluded

    def test_empty_valid_for_is_rejected(self):
        """An empty Valid For cell is a distinct error path from a malformed duration."""
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW)
        row[4] = ""  # Valid For
        ws = make_facts_sheet(wb, [row])
        facts, errors = parser.parse_facts(ws, ["CLS1"], ["phase1"])
        assert len(errors) == 1
        assert "Valid For" in errors[0]["error"] or "valid_for" in errors[0]["error"].lower()
