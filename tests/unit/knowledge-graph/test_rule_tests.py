#!/usr/bin/env python3
"""Tests for Sheet 10 (Rule Tests) parsing: the Facts cell's observation grammar
(parse_fact_observations) and the row-level sheet parser (parse_rule_tests). Spec §4.8.
"""
import importlib.util, os, pytest
from openpyxl import Workbook


def load_parser():
    script_path = os.path.join(os.path.dirname(__file__), "../../../scripts/generate-knowledge-graph.py")
    spec = importlib.util.spec_from_file_location("parser", os.path.abspath(script_path))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

parser = load_parser()


class TestParseFactObservations:
    def test_single_fact_defaults_to_now(self):
        obs, errs = parser.parse_fact_observations("systolic_bp = 185")
        assert errs == []
        assert obs == [{"fact": "systolic_bp", "value": 185.0, "offset_hours": 0}]

    def test_offset_is_negated_duration(self):
        obs, _ = parser.parse_fact_observations("weight = 70 @ -2d")
        assert obs[0]["offset_hours"] == -48

    def test_repeated_fact_builds_history(self):
        obs, _ = parser.parse_fact_observations("weight = 70 @ -2d\nweight = 72")
        assert len(obs) == 2
        assert obs[1]["offset_hours"] == 0

    def test_boolean_values(self):
        obs, _ = parser.parse_fact_observations("breathlessness = true\nedema = false")
        assert obs[0]["value"] is True and obs[1]["value"] is False

    def test_blank_lines_ignored(self):
        obs, _ = parser.parse_fact_observations("weight = 72\n\n  \nedema = true")
        assert len(obs) == 2

    def test_malformed_line_rejected(self):
        _, errs = parser.parse_fact_observations("weight 72")
        assert len(errs) == 1 and "weight 72" in errs[0]

    def test_positive_offset_rejected(self):
        _, errs = parser.parse_fact_observations("weight = 72 @ 2d")
        assert len(errs) == 1
        assert "must be negative" in errs[0].lower() or "past" in errs[0].lower()

    def test_malformed_offset_rejected(self):
        # NOTE: must be a value that MATCHES the regex but fails DURATION_RE.
        # "-2 days" fails the whole regex (the space breaks \S+), so it would
        # exercise the malformed-LINE branch and leave the malformed-OFFSET
        # branch untested while still passing a bare len(errs)==1 assertion.
        _, errs = parser.parse_fact_observations("weight = 72 @ -2x")
        assert len(errs) == 1
        assert "Malformed offset" in errs[0]

    def test_empty_value_rejected(self):
        _, errs = parser.parse_fact_observations("weight =")
        assert len(errs) == 1

    def test_double_equals_typo_rejected(self):
        # A plausible SME typo. Must NOT be silently accepted as the
        # categorical string "= 'low'" — categorical typechecking downstream
        # accepts any string, so nothing else would catch it.
        _, errs = parser.parse_fact_observations("mood == 'low'")
        assert len(errs) == 1

    def test_categorical_value(self):
        obs, errs = parser.parse_fact_observations("mood = 'irritable'\nmood2 = happy")
        assert errs == []
        assert obs[0]["value"] == "irritable"
        assert obs[1]["value"] == "happy"


# ---------------------------------------------------------------------------
# Sheet 10 row-level parser — parse_rule_tests(sheet, rules, facts)
# ---------------------------------------------------------------------------

FACTS = {
    "weight":         {"type": "number",  "unit": "kg", "display_name": "Weight",
                       "valid_for": "7d", "valid_for_hours": 168, "required": True,
                       "area": "vitals", "applicable_classifications": [],
                       "applicable_phases": [], "extraction_hint": None},
    "breathlessness": {"type": "boolean", "unit": None, "display_name": "Breathlessness",
                       "valid_for": "7d", "valid_for_hours": 168, "required": False,
                       "area": "symptom", "applicable_classifications": [],
                       "applicable_phases": [], "extraction_hint": None},
    "mood":           {"type": "categorical", "unit": None, "display_name": "Mood",
                       "valid_for": "7d", "valid_for_hours": 168, "required": False,
                       "area": "symptom", "applicable_classifications": [],
                       "applicable_phases": [], "extraction_hint": None},
}

RULES = [
    {"rule_id": "RULE001", "red_flag_id": "RF001", "order": 1,
     "expression": "weight >= 80", "ast": parser.parse_expression("weight >= 80"),
     "action": "ADVISE", "patient_action": None,
     "applicable_classifications": [], "applicable_phases": [],
     "reviewed_by": None, "reviewed_at": None},
]


def make_rule_tests_sheet(wb, rows):
    ws = wb.create_sheet("Rule Tests")
    ws.append(["Rule Tests"])
    ws.append(["Cases exercising each rule"])
    ws.append([])
    ws.append(["Test ID", "Rule ID", "Scenario", "Facts", "Expect", "Notes"])
    for r in rows:
        ws.append(r)
    return ws


# A valid base row
BASE_ROW = ["T_HF_001a", "RULE001", "Gained 2 kg in two days, now breathless",
            "weight = 70 @ -2d\nweight = 72\nbreathlessness = true",
            "FIRES", "Reviewed by Dr Smith"]


class TestParseRuleTests:
    def test_absent_sheet_returns_empty_triple(self):
        result = parser.parse_rule_tests(None, RULES, FACTS)
        assert result == ([], [], [])

    def test_valid_sheet_yields_case_with_populated_fields(self):
        wb = Workbook(); wb.remove(wb.active)
        make_rule_tests_sheet(wb, [BASE_ROW])
        cases, errors, warnings = parser.parse_rule_tests(wb["Rule Tests"], RULES, FACTS)
        assert errors == []
        assert warnings == []
        assert len(cases) == 1
        c = cases[0]
        assert c["test_id"] == "T_HF_001a"
        assert c["rule_id"] == "RULE001"
        assert c["scenario"] == "Gained 2 kg in two days, now breathless"
        assert c["expect"] == "FIRES"
        assert c["observations"] == [
            {"fact": "weight", "value": 70.0, "offset_hours": -48},
            {"fact": "weight", "value": 72.0, "offset_hours": 0},
            {"fact": "breathlessness", "value": True, "offset_hours": 0},
        ]
        assert c["notes"] == "Reviewed by Dr Smith"

    def test_unknown_rule_id_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[1] = "RULE999"
        make_rule_tests_sheet(wb, [row])
        cases, errors, _ = parser.parse_rule_tests(wb["Rule Tests"], RULES, FACTS)
        assert cases == []
        assert len(errors) >= 1
        assert any("RULE999" in e["error"] for e in errors)

    def test_unknown_fact_name_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[3] = "made_up_fact = 5"
        make_rule_tests_sheet(wb, [row])
        cases, errors, _ = parser.parse_rule_tests(wb["Rule Tests"], RULES, FACTS)
        assert cases == []
        assert len(errors) >= 1
        assert any("made_up_fact" in e["error"] for e in errors)

    def test_value_that_does_not_typecheck_is_rejected(self):
        # weight is declared numeric; 'true' is a boolean literal
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[3] = "weight = true"
        make_rule_tests_sheet(wb, [row])
        cases, errors, _ = parser.parse_rule_tests(wb["Rule Tests"], RULES, FACTS)
        assert cases == []
        assert len(errors) >= 1
        # Must be rejected as a typecheck failure against the declared type — not
        # misrouted through the unknown-fact branch (weight IS declared, just as
        # the wrong type here).
        assert any("weight" in e["error"] and "typecheck" in e["error"] and "number" in e["error"]
                   for e in errors)
        assert not any("unknown fact" in e["error"].lower() for e in errors)

    def test_invalid_expect_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[4] = "MAYBE"
        make_rule_tests_sheet(wb, [row])
        cases, errors, _ = parser.parse_rule_tests(wb["Rule Tests"], RULES, FACTS)
        assert cases == []
        assert len(errors) >= 1
        msg = errors[0]["error"]
        assert "FIRES" in msg and "NOT_FIRES" in msg

    def test_malformed_observation_line_carries_row_number(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[3] = "weight 72"  # missing '='
        make_rule_tests_sheet(wb, [row])
        ws = wb["Rule Tests"]
        cases, errors, _ = parser.parse_rule_tests(ws, RULES, FACTS)
        assert cases == []
        assert len(errors) >= 1
        # BASE_ROW is the only data row -> DATA_START_ROW (row 5)
        assert errors[0]["row"] == parser.DATA_START_ROW
        assert "weight 72" in errors[0]["error"]

    def test_notes_empty_is_accepted(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[5] = None
        make_rule_tests_sheet(wb, [row])
        cases, errors, _ = parser.parse_rule_tests(wb["Rule Tests"], RULES, FACTS)
        assert errors == []
        assert len(cases) == 1
        assert cases[0]["notes"] is None


# ---------------------------------------------------------------------------
# Coverage validation — validate_rule_test_coverage(rules, rule_tests). Spec §4.8.
# ---------------------------------------------------------------------------

TWO_RULES = [
    {"rule_id": "RULE001", "red_flag_id": "RF001", "order": 1},
    {"rule_id": "RULE002", "red_flag_id": "RF001", "order": 2},
]


def make_test_case(rule_id, expect):
    return {"rule_id": rule_id, "expect": expect}


class TestValidateRuleTestCoverage:
    def test_rule_with_both_fires_and_not_fires_has_no_warning(self):
        rules = [{"rule_id": "RULE001", "red_flag_id": "RF001", "order": 1}]
        rule_tests = [
            make_test_case("RULE001", "FIRES"),
            make_test_case("RULE001", "NOT_FIRES"),
        ]
        warnings = parser.validate_rule_test_coverage(rules, rule_tests)
        assert warnings == []

    def test_rule_with_only_fires_warns_about_missing_not_fires(self):
        rules = [{"rule_id": "RULE001", "red_flag_id": "RF001", "order": 1}]
        rule_tests = [make_test_case("RULE001", "FIRES")]
        warnings = parser.validate_rule_test_coverage(rules, rule_tests)
        assert len(warnings) == 1
        assert "RULE001" in warnings[0]
        assert "NOT_FIRES" in warnings[0]

    def test_rule_with_only_not_fires_warns_about_missing_fires(self):
        rules = [{"rule_id": "RULE001", "red_flag_id": "RF001", "order": 1}]
        rule_tests = [make_test_case("RULE001", "NOT_FIRES")]
        warnings = parser.validate_rule_test_coverage(rules, rule_tests)
        assert len(warnings) == 1
        assert "RULE001" in warnings[0]
        assert "FIRES" in warnings[0]

    def test_rule_with_no_cases_at_all_warns(self):
        rules = [{"rule_id": "RULE001", "red_flag_id": "RF001", "order": 1}]
        rule_tests = []
        warnings = parser.validate_rule_test_coverage(rules, rule_tests)
        assert len(warnings) == 1
        assert "RULE001" in warnings[0]
        assert "FIRES" in warnings[0] and "NOT_FIRES" in warnings[0]

    def test_no_rules_and_no_cases_produces_no_warnings(self):
        # An un-migrated KB with no Sheet 8 rules and no Sheet 10 cases must stay quiet.
        warnings = parser.validate_rule_test_coverage([], [])
        assert warnings == []

    def test_returns_plain_warning_strings(self):
        # validate_rule_test_coverage's contract: it returns a plain list of warning
        # strings, never dicts routed to `errors`. The end-to-end claim — that these
        # warnings never flip a KB's validity through parse_xlsx — is covered by
        # TestParseXlsxRuleTestsWiring below.
        rules = [{"rule_id": "RULE001", "red_flag_id": "RF001", "order": 1}]
        warnings = parser.validate_rule_test_coverage(rules, [])
        assert len(warnings) == 1
        assert all(isinstance(w, str) for w in warnings)

    def test_two_rules_only_missing_one_reports_only_that_one(self):
        rule_tests = [
            make_test_case("RULE001", "FIRES"),
            make_test_case("RULE001", "NOT_FIRES"),
            make_test_case("RULE002", "FIRES"),
        ]
        warnings = parser.validate_rule_test_coverage(TWO_RULES, rule_tests)
        assert len(warnings) == 1
        assert "RULE002" in warnings[0]


# ---------------------------------------------------------------------------
# Integration: parse_xlsx() end-to-end wiring for Sheet 10. Spec §4.8.
#
# Every test above calls parse_rule_tests()/validate_rule_test_coverage() directly.
# This is the one test that drives the real entry point — the workbook is built
# fresh here (not via test_parser.py's make_minimal_workbook, which is quarantined:
# it targets the old singular "Condition & Phases" sheet name and predates the
# Facts/Rules/Rule Tests sheets entirely) so a typo'd sheet name, a dropped line,
# or a swapped argument in a future parse_xlsx refactor has something to trip.
# ---------------------------------------------------------------------------

def _append_titled_sheet(wb, name, headers, data_rows=()):
    """Build a sheet matching the fixed layout every parser here assumes:
    row 1 title, row 2 description, row 3 blank, row 4 header (HEADER_ROW),
    row 5+ data (DATA_START_ROW)."""
    ws = wb.create_sheet(name)
    ws.append([name])
    ws.append([f"{name} — test fixture"])
    ws.append([])
    ws.append(headers)
    for row in data_rows:
        ws.append(row)
    return ws


def make_full_valid_workbook_with_uncovered_rule():
    """A minimal workbook satisfying every required tab, plus Facts/Rules/Rule
    Tests, where the single rule has a FIRES case but no NOT_FIRES case."""
    wb = Workbook(); wb.remove(wb.active)

    _append_titled_sheet(wb, "Conditions & Phases", [
        "Condition", "Condition Type", "Trigger", "Phase Model",
        "Classification", "Phase", "Day Range", "Focus", "Review Point",
        "Sheet Version",
    ], [
        ["Test Condition", "episodic", "surgery_date", "linear",
         "TEST_CLS", "Phase I", "0-14", "Recovery focus", "Review wound", "v1"],
    ])
    _append_titled_sheet(wb, "Symptoms", [
        "Symptom ID", "Symptom Name", "Applicable Classifications",
        "Applicable Phases", "Base Severity", "Severity Score",
        "Phase Override?", "Override Detail", "Clinical Notes",
    ])
    _append_titled_sheet(wb, "Assessment Questions", [
        "Question ID", "Symptom ID", "Order", "Question Text (voice prompt)",
        "Answer A (label)", "Risk Shift A", "Escalate? A",
        "Answer B (label)", "Risk Shift B", "Escalate? B", "Next Q",
    ])
    _append_titled_sheet(wb, "Red Flags", [
        "Red Flag ID", "Symptom ID", "Trigger Condition",
        "Applicable Classifications", "Applicable Phases",
        "Action", "Urgency", "Context Note", "Clinical Rationale",
    ], [
        ["RF001", "—", "Weight gain of 2kg in 48h", "ALL", "ALL",
         "ESCALATE", "immediate", "", "Fluid overload risk"],
    ])
    _append_titled_sheet(wb, "Instructions", [
        "Instruction ID", "Applicable Classifications", "Phase",
        "Category", "Instruction Text", "Issued At", "Replaces", "Notes",
    ])
    _append_titled_sheet(wb, "Scoring Logic", ["Key", "Value"])
    _append_titled_sheet(wb, "Facts", [
        "Machine Name", "Display Name", "Type", "Unit",
        "Valid For", "Required", "Area",
        "Applicable Classifications", "Applicable Phases", "Extraction Hint",
    ], [
        ["weight", "Weight", "number", "kg", "7d", "Yes", "vitals",
         "ALL", "ALL", "Patient weight"],
    ])
    _append_titled_sheet(wb, "Rules", [
        "Rule ID", "Red Flag ID", "Order", "Expression",
        "Action", "Patient Action",
        "Applicable Classifications", "Applicable Phases",
        "Reviewed By", "Reviewed At",
    ], [
        ["RULE001", "RF001", 1, "weight >= 80",
         "ADVISE", "SELF_MONITOR", "ALL", "ALL", "Dr Smith", "2026-07-01"],
    ])
    make_rule_tests_sheet(wb, [
        # Only a FIRES case — deliberately no NOT_FIRES case for RULE001.
        ["T001", "RULE001", "Weight above threshold", "weight = 90", "FIRES", ""],
    ])
    return wb


class TestParseXlsxRuleTestsWiring:
    def test_parse_xlsx_wires_rule_tests_and_emits_coverage_warning(self, tmp_path):
        wb = make_full_valid_workbook_with_uncovered_rule()
        path = tmp_path / "kb.xlsx"
        wb.save(str(path))

        result = parser.parse_xlsx(str(path))

        # A future typo'd sheet name / dropped call / swapped argument in
        # parse_xlsx would surface here as a structural error or missing key.
        assert result["errors"] == []

        kg = result["knowledge_graph"]
        assert kg is not None
        assert kg["rule_tests"] == [
            {
                "test_id": "T001",
                "rule_id": "RULE001",
                "scenario": "Weight above threshold",
                "observations": [{"fact": "weight", "value": 90.0, "offset_hours": 0}],
                "expect": "FIRES",
                "notes": None,
            }
        ]

        assert any("RULE001" in w and "NOT_FIRES" in w for w in result["warnings"])

        # The coverage warning must never demote validity — that's the whole point
        # of it being a warning, not an error.
        assert result["valid"] is True
        assert result["errors"] == []
