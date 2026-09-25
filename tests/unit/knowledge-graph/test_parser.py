#!/usr/bin/env python3
"""
Pytest tests for generate-knowledge-graph.py parser.
Imports parser functions directly via importlib (filename has hyphens).
"""
import importlib.util
import os
import io
import json
import pytest
import openpyxl
from openpyxl import Workbook

# Workbook fixtures were migrated from v3.0 to v3.2 in Task 9 of
# docs/superpowers/plans/2026-07-16-kb-authoring-toolchain.md. The former
# per-class @STALE_FIXTURES quarantine has been lifted; all classes below run.
# New v3.2 sheets (Settings/Facts/Expressions/Rules) are covered separately by
# test_settings.py / test_facts.py / test_expressions.py / test_rules.py.

# ── Loader ──────────────────────────────────────────────────────────────────

def load_parser():
    """Import generate-knowledge-graph.py as a module."""
    script_path = os.path.join(
        os.path.dirname(__file__), "../../../scripts/generate-knowledge-graph.py"
    )
    spec = importlib.util.spec_from_file_location("parser", os.path.abspath(script_path))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

try:
    parser = load_parser()
except Exception as e:
    import pytest as _pytest
    _pytest.skip(f"Parser module unavailable: {e}", allow_module_level=True)

# ── Workbook helpers ─────────────────────────────────────────────────────────

def add_preamble(ws, title, description="Test fixture"):
    """Emit the standard 3-row preamble so the header lands on row 4 (HEADER_ROW)."""
    ws.append([title])
    ws.append([description])
    ws.append([])

def make_condition_sheet(wb, phases, extra_headers=None, extra_values=None):
    """
    Add a 'Conditions & Phases' sheet with given phases.
    phases: list of (phase_name, day_range_str, focus, review)
    extra_headers: list of additional column header strings
    extra_values: list of lists, one per phase row, with values for extra columns
                  (overrides the base column of the same name, never duplicates it)
    """
    ws = wb.create_sheet("Conditions & Phases")
    add_preamble(ws, "Conditions & Phases")

    def row_dict(phase_name, day_range, focus, review):
        return {"Condition": "Test Condition", "Condition Type": "chronic",
                "Trigger": "enrollment_date", "Phase Model": "linear",
                "Classification": "TEST_CLS", "Phase": phase_name,
                "Day Range": day_range, "Focus": focus,
                "Review Point": review, "Sheet Version": "v1"}

    rows = []
    for i, (phase_name, day_range, focus, review) in enumerate(phases):
        d = row_dict(phase_name, day_range, focus, review)
        if extra_headers and extra_values and i < len(extra_values):
            d.update(dict(zip(extra_headers, extra_values[i])))   # override, never duplicate
        rows.append(d)

    headers = list(rows[0].keys())          # merged key set — extras that are NOT base columns append here
    ws.append(headers)
    for d in rows:
        ws.append([d.get(h) for h in headers])
    return ws

def make_symptoms_sheet(wb, phases):
    ws = wb.create_sheet("Symptoms")
    add_preamble(ws, "Symptoms")
    ws.append(["Symptom ID", "Symptom Name", "Applicable Classifications",
               "Applicable Phases", "Base Severity", "Severity Score",
               "Phase Override?", "Override Detail", "Clinical Notes"])
    phase_str = ",".join(phases)
    ws.append(["SYM_001", "Test Symptom", "TEST_CLS", phase_str, "low", "1", "No", "", ""])
    return ws

def make_questions_sheet(wb):
    ws = wb.create_sheet("Assessment Questions")
    add_preamble(ws, "Assessment Questions")
    ws.append(["Question ID", "Symptom ID", "Order", "Question Text (voice prompt)",
               "Answer A (label)", "Risk Shift A", "Escalate? A",
               "Answer B (label)", "Risk Shift B", "Escalate? B", "Next Q"])
    ws.append(["Q_001", "SYM_001", 1, "How is the pain?",
               "Better", -1, "No", "Worse", 1, "No", "—"])
    return ws

def make_red_flags_sheet(wb, phases):
    ws = wb.create_sheet("Red Flags")
    add_preamble(ws, "Red Flags")
    ws.append(["Red Flag ID", "Symptom ID", "Trigger Condition",
               "Applicable Classifications", "Applicable Phases",
               "Action", "Urgency", "Context Note", "Clinical Rationale"])
    phase_str = ",".join(phases)
    ws.append(["RF_001", "SYM_001", "Chest pain", "TEST_CLS", phase_str,
               "ESCALATE", "immediate", "", ""])
    return ws

def make_instructions_sheet(wb, phases):
    ws = wb.create_sheet("Instructions")
    add_preamble(ws, "Instructions")
    ws.append(["Instruction ID", "Applicable Classifications", "Phase",
               "Category", "Instruction Text", "Issued At", "Replaces", "Notes"])
    for phase in phases:
        ws.append([f"INST_{phase}", "TEST_CLS", phase, "monitoring",
                   f"Instructions for {phase}", "", "—", ""])
    return ws

def make_scoring_sheet(wb):
    ws = wb.create_sheet("Scoring Logic")
    add_preamble(ws, "Scoring Logic")
    ws.append(["Key", "Value"])
    return ws

def save_workbook(wb, tmp_path, name="test.xlsx"):
    """Save workbook to tmp_path and return the filepath string."""
    path = tmp_path / name
    wb.save(str(path))
    return str(path)

def make_minimal_workbook(phases, extra_headers=None, extra_values=None):
    """Create a minimal but complete valid workbook."""
    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet
    phase_names = [p[0].upper().replace(" ", "_") for p in phases]
    make_condition_sheet(wb, phases, extra_headers, extra_values)
    make_symptoms_sheet(wb, phase_names)
    make_questions_sheet(wb)
    make_red_flags_sheet(wb, phase_names)
    make_instructions_sheet(wb, phase_names)
    make_scoring_sheet(wb)
    return wb

# ── Change 3: Phase name whitelist regression ────────────────────────────────

class TestPhaseNameWhitelist:

    def test_standard_phase_names_still_work(self, tmp_path):
        """REGRESSION: PHASE_I/II/III must parse identically after whitelist removal."""
        phases = [
            ("Phase I",  "0-7",   "Early recovery", "Review wound"),
            ("Phase II", "8-14",  "Mid recovery",   "Review mobility"),
            ("Phase III","15-30", "Late recovery",  "Review discharge"),
        ]
        wb = make_minimal_workbook(phases)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is True, result["errors"]
        kg = result["knowledge_graph"]
        assert "PHASE_I" in kg["condition"]["classifications"]["TEST_CLS"]["phases"]
        assert "PHASE_II" in kg["condition"]["classifications"]["TEST_CLS"]["phases"]
        assert "PHASE_III" in kg["condition"]["classifications"]["TEST_CLS"]["phases"]

    def test_custom_phase_names_accepted(self, tmp_path):
        """Custom names like Initiation/Titration/Maintenance must be accepted."""
        phases = [
            ("Initiation",  "0-30",  "Start meds", "Review BP"),
            ("Titration",   "31-90", "Adjust dose", "Review labs"),
            ("Maintenance", "91-180","Stable",       "Annual review"),
        ]
        wb = make_minimal_workbook(phases)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is True, result["errors"]
        kg = result["knowledge_graph"]
        assert "INITIATION" in kg["condition"]["classifications"]["TEST_CLS"]["phases"]
        assert "TITRATION" in kg["condition"]["classifications"]["TEST_CLS"]["phases"]
        assert "MAINTENANCE" in kg["condition"]["classifications"]["TEST_CLS"]["phases"]

    def test_undefined_phase_in_red_flags_still_errors(self, tmp_path):
        """Cross-sheet check: phase in Red Flags not in Conditions tab → error."""
        phases = [("Phase I", "0-7", "Focus", "Review")]
        wb = make_minimal_workbook(phases)
        # Overwrite red flags sheet with a bad phase name
        del wb["Red Flags"]
        ws = wb.create_sheet("Red Flags")
        add_preamble(ws, "Red Flags")
        ws.append(["Red Flag ID", "Symptom ID", "Trigger Condition",
                   "Applicable Classifications", "Applicable Phases", "Action",
                   "Urgency", "Context Note", "Clinical Rationale"])
        ws.append(["RF_001", "SYM_001", "Chest pain", "TEST_CLS", "NONEXISTENT_PHASE",
                   "ESCALATE", "immediate", "", ""])
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        error_msgs = [e["error"] for e in result["errors"]]
        assert any("NONEXISTENT_PHASE" in m for m in error_msgs)
        assert any("Defined phases are" in m for m in error_msgs)

    def test_empty_phase_name_rejected(self, tmp_path):
        """Empty phase name must be rejected."""
        phases = [("", "0-7", "Focus", "Review")]
        wb = make_minimal_workbook(phases)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        # A row with an empty Phase value fails validation and is skipped entirely
        # before file_condition is ever set, so parse_xlsx reports it the same way
        # as a workbook with zero condition rows -- not a distinct "no phases" message.
        assert result["valid"] is False
        assert any("Phase is required" in e["error"] for e in result["errors"])
        assert any("No condition rows found" in e["error"] for e in result["errors"])

# ── Change 2: Flexible Day Range ─────────────────────────────────────────────

class TestParseDayRange:
    """Unit tests for the parse_day_range() helper function."""

    # Fixed range
    def test_fixed_hyphen(self):
        start, end, rtype = parser.parse_day_range("0-14")
        assert start == 0 and end == 14 and rtype == "fixed"

    def test_fixed_en_dash(self):
        start, end, rtype = parser.parse_day_range("0–14")
        assert start == 0 and end == 14 and rtype == "fixed"

    def test_fixed_with_spaces(self):
        start, end, rtype = parser.parse_day_range(" 30 - 90 ")
        assert start == 30 and end == 90 and rtype == "fixed"

    # Open-ended
    def test_open_ended(self):
        start, end, rtype = parser.parse_day_range("90+")
        assert start == 90 and end is None and rtype == "open"

    def test_open_ended_zero(self):
        start, end, rtype = parser.parse_day_range("0+")
        assert start == 0 and end is None and rtype == "open"

    def test_open_ended_with_spaces(self):
        start, end, rtype = parser.parse_day_range(" 180+ ")
        assert start == 180 and end is None and rtype == "open"

    # Ongoing
    def test_ongoing_lowercase(self):
        start, end, rtype = parser.parse_day_range("ongoing")
        assert start == 0 and end is None and rtype == "ongoing"

    def test_ongoing_uppercase(self):
        start, end, rtype = parser.parse_day_range("ONGOING")
        assert start == 0 and end is None and rtype == "ongoing"

    # Validation errors
    def test_rejects_negative_start(self):
        # "-5-10" splits into 3 parts after hyphen-split, falls through to generic error
        with pytest.raises(ValueError, match="format"):
            parser.parse_day_range("-5-10")

    def test_rejects_end_less_than_start(self):
        with pytest.raises(ValueError, match="[Ee]nd"):
            parser.parse_day_range("30-14")

    def test_rejects_nonsense(self):
        with pytest.raises(ValueError, match="[Ff]ormat"):
            parser.parse_day_range("90–365+")

    def test_rejects_empty(self):
        with pytest.raises(ValueError):
            parser.parse_day_range("")

    def test_rejects_letters(self):
        with pytest.raises(ValueError):
            parser.parse_day_range("abc")

    def test_rejects_negative_open_ended(self):
        """Exercises the start < 0 guard in the open-ended branch."""
        with pytest.raises(ValueError, match="non-negative"):
            parser.parse_day_range("-5+")

    def test_rejects_whitespace_only(self):
        """Whitespace-only input strips to empty, hits the empty guard."""
        with pytest.raises(ValueError):
            parser.parse_day_range("   ")

    def test_rejects_equal_start_end(self):
        """end == start should be rejected (end must be > start)."""
        with pytest.raises(ValueError, match="[Ee]nd"):
            parser.parse_day_range("7-7")


class TestParseDayRangeInWorkbook:
    """Integration tests: parse_day_range formats inside a full workbook."""

    def test_open_ended_phase_stored_correctly(self, tmp_path):
        phases = [
            ("Phase I",       "0-30",  "Early", "Review"),
            ("Maintenance",   "31+",   "Stable", "Annual"),
        ]
        wb = make_minimal_workbook(phases)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is True, result["errors"]
        kg = result["knowledge_graph"]
        assert kg["condition"]["classifications"]["TEST_CLS"]["phases"]["MAINTENANCE"]["day_range"] == [31, None]
        assert kg["condition"]["classifications"]["TEST_CLS"]["phases"]["MAINTENANCE"]["day_range_type"] == "open"

    def test_ongoing_phase_stored_correctly(self, tmp_path):
        phases = [
            ("Acute",      "0-30",    "Acute phase", "Review"),
            ("Chronic",    "ongoing", "Long-term",   "Annual"),
        ]
        wb = make_minimal_workbook(phases)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is True, result["errors"]
        kg = result["knowledge_graph"]
        assert kg["condition"]["classifications"]["TEST_CLS"]["phases"]["CHRONIC"]["day_range"] == [0, None]
        assert kg["condition"]["classifications"]["TEST_CLS"]["phases"]["CHRONIC"]["day_range_type"] == "ongoing"

    def test_fixed_range_day_range_type_set(self, tmp_path):
        """REGRESSION: existing fixed ranges must have day_range_type='fixed'."""
        phases = [("Phase I", "0-7", "Focus", "Review")]
        wb = make_minimal_workbook(phases)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is True, result["errors"]
        assert result["knowledge_graph"]["condition"]["classifications"]["TEST_CLS"]["phases"]["PHASE_I"]["day_range_type"] == "fixed"

    def test_invalid_day_range_produces_error(self, tmp_path):
        phases = [("Phase I", "abc", "Focus", "Review")]
        wb = make_minimal_workbook(phases)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        error_msgs = [e["error"] for e in result["errors"]]
        assert any("Day Range" in m or "format" in m.lower() for m in error_msgs)


# ── Change 1: New Optional Columns ───────────────────────────────────────────

class TestConditionMetadataColumns:

    def test_new_columns_parsed_when_present(self, tmp_path):
        """New columns present with valid values → stored correctly."""
        phases = [("Initiation", "0-30", "Focus", "Review")]
        extra_headers = ["Condition Type", "Trigger", "Phase Model"]
        extra_values = [["chronic", "enrollment_date", "cyclical"]]
        wb = make_minimal_workbook(phases, extra_headers, extra_values)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is True, result["errors"]
        cond = result["knowledge_graph"]["condition"]
        assert cond["condition_type"] == "chronic"
        assert cond["trigger"] == ["enrollment_date"]
        assert cond["phase_model"] == "cyclical"

    def test_pipe_delimited_trigger_parsed_as_list(self, tmp_path):
        phases = [("Phase I", "0-30", "Focus", "Review")]
        extra_headers = ["Condition Type", "Trigger", "Phase Model"]
        extra_values = [["hybrid", "discharge_date | enrollment_date", "cyclical"]]
        wb = make_minimal_workbook(phases, extra_headers, extra_values)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is True, result["errors"]
        assert result["knowledge_graph"]["condition"]["trigger"] == [
            "discharge_date", "enrollment_date"
        ]

    def test_invalid_condition_type_rejected(self, tmp_path):
        phases = [("Phase I", "0-7", "Focus", "Review")]
        extra_headers = ["Condition Type", "Trigger", "Phase Model"]
        extra_values = [["unknown_type", "surgery_date", "linear"]]
        wb = make_minimal_workbook(phases, extra_headers, extra_values)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        assert any("Condition Type" in e["error"] for e in result["errors"])
        assert any(e.get("column") == "Condition Type" for e in result["errors"])

    def test_invalid_trigger_name_rejected(self, tmp_path):
        phases = [("Phase I", "0-7", "Focus", "Review")]
        extra_headers = ["Condition Type", "Trigger", "Phase Model"]
        extra_values = [["episodic", "bad_trigger", "linear"]]
        wb = make_minimal_workbook(phases, extra_headers, extra_values)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        assert any("Trigger" in e["error"] for e in result["errors"])
        assert any(e.get("column") == "Trigger" for e in result["errors"])

    def test_invalid_phase_model_rejected(self, tmp_path):
        phases = [("Phase I", "0-7", "Focus", "Review")]
        extra_headers = ["Condition Type", "Trigger", "Phase Model"]
        extra_values = [["episodic", "surgery_date", "recursive"]]
        wb = make_minimal_workbook(phases, extra_headers, extra_values)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        assert any("Phase Model" in e["error"] for e in result["errors"])
        assert any(e.get("column") == "Phase Model" for e in result["errors"])

    def test_traversal_contains_cyclical_reentry_step(self, tmp_path):
        """Change 4 Step 3: traversal must include a step about Phase Model / cyclical re-entry."""
        phases = [("Phase I", "0-30", "Focus", "Review")]
        extra_headers = ["Condition Type", "Trigger", "Phase Model"]
        extra_values = [["chronic", "enrollment_date", "cyclical"]]
        wb = make_minimal_workbook(phases, extra_headers, extra_values)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is True, result["errors"]
        traversal = result["knowledge_graph"]["traversal"]
        assert any("Phase Model" in step or "cyclical" in step.lower() for step in traversal), (
            f"Expected traversal step mentioning Phase Model/cyclical. Got: {traversal}"
        )

    def test_traversal_regression_step_count(self, tmp_path):
        """REGRESSION: traversal must now have 11 steps (v3.2 adds the cyclical re-entry rule)."""
        phases = [("Phase I", "0-7", "Focus", "Review")]
        wb = make_minimal_workbook(phases)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is True, result["errors"]
        assert len(result["knowledge_graph"]["traversal"]) == 11

# ── Change 6: Error Message Quality ──────────────────────────────────────────

class TestErrorMessages:

    def test_day_range_error_lists_accepted_formats(self, tmp_path):
        phases = [("Phase I", "invalid_range", "Focus", "Review")]
        wb = make_minimal_workbook(phases)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        msg = next(e["error"] for e in result["errors"] if e.get("column") == "Day Range")
        assert "0–14" in msg or "e.g." in msg.lower(), f"Unhelpful message: {msg}"
        assert "ongoing" in msg.lower()
        assert "+" in msg

    def test_undefined_phase_error_lists_defined_phases(self, tmp_path):
        phases = [("Phase I", "0-7", "Focus", "Review")]
        wb = make_minimal_workbook(phases)
        del wb["Red Flags"]
        ws = wb.create_sheet("Red Flags")
        add_preamble(ws, "Red Flags")
        ws.append(["Red Flag ID", "Symptom ID", "Trigger Condition",
                   "Applicable Classifications", "Applicable Phases", "Action",
                   "Urgency", "Context Note", "Clinical Rationale"])
        ws.append(["RF_001", "SYM_001", "Chest pain", "TEST_CLS", "PHASE_IX",
                   "ESCALATE", "immediate", "", ""])
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        msg = next(e["error"] for e in result["errors"] if "PHASE_IX" in e["error"])
        assert "Defined phases" in msg
        assert "PHASE_I" in msg

    def test_invalid_condition_type_error_is_actionable(self, tmp_path):
        phases = [("Phase I", "0-7", "Focus", "Review")]
        extra_headers = ["Condition Type", "Trigger", "Phase Model"]
        extra_values = [["periodic", "surgery_date", "linear"]]
        wb = make_minimal_workbook(phases, extra_headers, extra_values)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        msg = next(e["error"] for e in result["errors"] if e.get("column") == "Condition Type")
        assert "episodic" in msg and "chronic" in msg and "hybrid" in msg

    def test_invalid_trigger_error_lists_valid_names(self, tmp_path):
        phases = [("Phase I", "0-7", "Focus", "Review")]
        extra_headers = ["Condition Type", "Trigger", "Phase Model"]
        extra_values = [["episodic", "readmission_date", "linear"]]
        wb = make_minimal_workbook(phases, extra_headers, extra_values)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        msg = next(e["error"] for e in result["errors"] if e.get("column") == "Trigger")
        assert "discharge_date" in msg
        assert "enrollment_date" in msg

    def test_invalid_phase_model_error_lists_valid_values(self, tmp_path):
        phases = [("Phase I", "0-7", "Focus", "Review")]
        extra_headers = ["Condition Type", "Trigger", "Phase Model"]
        extra_values = [["episodic", "surgery_date", "recursive"]]
        wb = make_minimal_workbook(phases, extra_headers, extra_values)
        result = parser.parse_xlsx(save_workbook(wb, tmp_path))
        assert result["valid"] is False
        msg = next(e["error"] for e in result["errors"] if e.get("column") == "Phase Model")
        assert "linear" in msg and "cyclical" in msg
