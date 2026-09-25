#!/usr/bin/env python3
"""Tests for Sheet 9 (Settings) parsing. Spec §4.7."""
import importlib.util, os, pytest
from openpyxl import Workbook


def load_parser():
    script_path = os.path.join(os.path.dirname(__file__), "../../../scripts/generate-knowledge-graph.py")
    spec = importlib.util.spec_from_file_location("parser", os.path.abspath(script_path))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

parser = load_parser()


def make_settings_sheet(wb, rows):
    ws = wb.create_sheet("Settings")
    ws.append(["Settings"])
    ws.append(["Condition-level configuration"])
    ws.append([])
    ws.append(["Setting", "Value"])
    for k, v in rows:
        ws.append([k, v])
    return ws


class TestSettings:
    def test_defaults_when_sheet_absent(self):
        settings, errors, warnings = parser.parse_settings(None)
        assert errors == [] and warnings == []
        assert settings["max_questions_per_checkin"] == 6
        assert settings["min_days_between_checkins"] == "3d"
        assert settings["checkin_trigger_stale_count"] == 1
        assert settings["time_uncertainty_tolerance"] == 0.25

    def test_explicit_values_override_defaults(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_settings_sheet(wb, [("max_questions_per_checkin", 4),
                                      ("time_uncertainty_tolerance", 0.5)])
        settings, errors, warnings = parser.parse_settings(ws)
        assert errors == []
        assert settings["max_questions_per_checkin"] == 4
        assert settings["time_uncertainty_tolerance"] == 0.5
        # untouched settings keep defaults
        assert settings["checkin_trigger_stale_count"] == 1

    def test_unknown_setting_warns_but_does_not_fail(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_settings_sheet(wb, [("future_setting", "x")])
        settings, errors, warnings = parser.parse_settings(ws)
        assert errors == []                                   # forward compatible
        assert "future_setting" not in settings
        assert any("future_setting" in w for w in warnings)   # but it must SAY so

    def test_malformed_value_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_settings_sheet(wb, [("max_questions_per_checkin", "not_a_number")])
        settings, errors, warnings = parser.parse_settings(ws)
        assert len(errors) == 1
        assert "max_questions_per_checkin" in errors[0]["error"]

    def test_malformed_duration_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_settings_sheet(wb, [("min_days_between_checkins", "3 days")])
        settings, errors, warnings = parser.parse_settings(ws)
        assert len(errors) == 1
        assert "min_days_between_checkins" in errors[0]["error"]

    def test_missing_column_header_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = wb.create_sheet("Settings")
        ws.append(["Settings"])
        ws.append(["Config"])
        ws.append([])
        ws.append(["BadCol", "OtherCol"])  # "Setting" column absent
        settings, errors, warnings = parser.parse_settings(ws)
        assert len(errors) == 1
        assert "Setting" in errors[0]["error"] or "Value" in errors[0]["error"]

    def test_valid_duration_is_accepted(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_settings_sheet(wb, [("min_days_between_checkins", "7d")])
        settings, errors, warnings = parser.parse_settings(ws)
        assert errors == []
        assert settings["min_days_between_checkins"] == "7d"

    def test_string_setting_accepted(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_settings_sheet(wb, [("escalation_message", "Call 080-66202020 to book an appointment.")])
        settings, errors, warnings = parser.parse_settings(ws)
        assert errors == []
        assert settings["escalation_message"] == "Call 080-66202020 to book an appointment."

    def test_string_setting_trims_whitespace(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_settings_sheet(wb, [("escalation_message", "  padded value  ")])
        settings, errors, warnings = parser.parse_settings(ws)
        assert errors == []
        assert settings["escalation_message"] == "padded value"

    def test_string_setting_placeholder_is_unset_not_error(self):
        wb = Workbook(); wb.remove(wb.active)
        ws = make_settings_sheet(wb, [("escalation_message", "-")])
        settings, errors, warnings = parser.parse_settings(ws)
        assert errors == []
        assert settings["escalation_message"] is None

    def test_string_setting_defaults_to_none_when_absent(self):
        settings, errors, warnings = parser.parse_settings(None)
        assert settings["escalation_message"] is None
