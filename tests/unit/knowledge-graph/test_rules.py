#!/usr/bin/env python3
"""Tests for Sheet 8 (Rules) parsing. Spec §4.6."""
import importlib.util, os, pytest
from openpyxl import Workbook


def load_parser():
    script_path = os.path.join(os.path.dirname(__file__), "../../../scripts/generate-knowledge-graph.py")
    spec = importlib.util.spec_from_file_location("parser", os.path.abspath(script_path))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

parser = load_parser()


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

KNOWN_CLS = {"TEST_CLS"}
KNOWN_PHASES = {"phase1"}

FACTS = {
    "weight":        {"type": "number",  "unit": "kg", "display_name": "Weight",
                      "valid_for": "7d", "valid_for_hours": 168, "required": True,
                      "area": "vitals", "applicable_classifications": [],
                      "applicable_phases": [], "extraction_hint": None},
    "breathlessness": {"type": "boolean", "unit": None, "display_name": "Breathlessness",
                       "valid_for": "7d", "valid_for_hours": 168, "required": False,
                       "area": "symptom", "applicable_classifications": [],
                       "applicable_phases": [], "extraction_hint": None},
}

RED_FLAGS = {"RF001": {"red_flag_id": "RF001"}}


def make_rules_sheet(wb, rows):
    ws = wb.create_sheet("Rules")
    ws.append(["Rules"])
    ws.append(["Clinical decision rules"])
    ws.append([])
    ws.append(["Rule ID", "Red Flag ID", "Order", "Expression",
               "Action", "Patient Action",
               "Applicable Classifications", "Applicable Phases",
               "Reviewed By", "Reviewed At"])
    for r in rows:
        ws.append(r)
    return ws


# A valid base row
BASE_ROW = ["RULE001", "RF001", 1,
            "weight >= 80",
            "ADVISE", "SELF_MONITOR",
            "TEST_CLS", "phase1",
            "Dr Smith", "2026-07-01"]


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestRules:
    def test_absent_sheet_returns_empty_list(self):
        rules, errors = parser.parse_rules(None, FACTS, RED_FLAGS, KNOWN_CLS, KNOWN_PHASES)
        assert rules == [] and errors == []

    def test_valid_sheet_yields_rules_with_ast(self):
        wb = Workbook(); wb.remove(wb.active)
        make_rules_sheet(wb, [BASE_ROW])
        rules, errors = parser.parse_rules(wb["Rules"], FACTS, RED_FLAGS, KNOWN_CLS, KNOWN_PHASES)
        assert errors == []
        assert len(rules) == 1
        r = rules[0]
        assert r["rule_id"] == "RULE001"
        assert r["red_flag_id"] == "RF001"
        assert r["order"] == 1
        assert r["expression"] == "weight >= 80"
        assert r["ast"]["kind"] == "compare"
        assert r["action"] == "ADVISE"
        assert r["patient_action"] == "SELF_MONITOR"
        assert r["reviewed_by"] == "Dr Smith"
        assert r["reviewed_at"] == "2026-07-01"

    def test_red_flag_id_not_in_sheet4_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[1] = "RF999"  # unknown
        make_rules_sheet(wb, [row])
        rules, errors = parser.parse_rules(wb["Rules"], FACTS, RED_FLAGS, KNOWN_CLS, KNOWN_PHASES)
        assert len(errors) >= 1
        assert any("RF999" in e["error"] for e in errors)

    def test_unparseable_expression_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[3] = "weight $ 80"  # invalid char
        make_rules_sheet(wb, [row])
        rules, errors = parser.parse_rules(wb["Rules"], FACTS, RED_FLAGS, KNOWN_CLS, KNOWN_PHASES)
        assert len(errors) >= 1
        assert any("RULE001" in e["error"] for e in errors)

    def test_expression_referencing_undeclared_fact_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[3] = "undeclared_fact >= 5"
        make_rules_sheet(wb, [row])
        rules, errors = parser.parse_rules(wb["Rules"], FACTS, RED_FLAGS, KNOWN_CLS, KNOWN_PHASES)
        assert len(errors) >= 1
        assert any("undeclared_fact" in e["error"] for e in errors)

    def test_invalid_action_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[4] = "INVALID"
        make_rules_sheet(wb, [row])
        rules, errors = parser.parse_rules(wb["Rules"], FACTS, RED_FLAGS, KNOWN_CLS, KNOWN_PHASES)
        assert len(errors) >= 1
        assert any("Action" in e["error"] or "INVALID" in e["error"] for e in errors)

    def test_invalid_patient_action_is_rejected(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[5] = "do_something_weird"
        make_rules_sheet(wb, [row])
        rules, errors = parser.parse_rules(wb["Rules"], FACTS, RED_FLAGS, KNOWN_CLS, KNOWN_PHASES)
        assert len(errors) >= 1

    def test_invalid_applicable_classification_produces_readable_error(self):
        """Error message must contain the unknown value, not a raw dict repr (regression guard)."""
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[6] = "UNKNOWN_CLS"
        make_rules_sheet(wb, [row])
        _, errors = parser.parse_rules(wb["Rules"], FACTS, RED_FLAGS, KNOWN_CLS, KNOWN_PHASES)
        assert len(errors) >= 1
        msg = errors[0]["error"]
        assert "UNKNOWN_CLS" in msg
        assert "{" not in msg, f"Error message contains raw dict repr: {msg!r}"

    def test_reviewed_by_and_reviewed_at_empty_is_accepted(self):
        wb = Workbook(); wb.remove(wb.active)
        row = list(BASE_ROW); row[8] = None; row[9] = None
        make_rules_sheet(wb, [row])
        rules, errors = parser.parse_rules(wb["Rules"], FACTS, RED_FLAGS, KNOWN_CLS, KNOWN_PHASES)
        assert errors == []
        assert rules[0]["reviewed_by"] is None
        assert rules[0]["reviewed_at"] is None


class TestRuleGroupValidation:
    """Tests for validate_rule_groups(rules, facts, settings). Spec §5.5 cross-rule checks."""

    def _settings(self, cadence="3d"):
        return {
            "max_questions_per_checkin": 6,
            "min_days_between_checkins": cadence,
            "checkin_trigger_stale_count": 1,
            "time_uncertainty_tolerance": 0.25,
        }

    def _rule(self, rule_id, red_flag_id, order, expression, action="ADVISE"):
        return {
            "rule_id": rule_id, "red_flag_id": red_flag_id, "order": order,
            "expression": expression,
            "ast": parser.parse_expression(expression),
            "action": action, "patient_action": None,
            "applicable_classifications": [], "applicable_phases": [],
            "reviewed_by": None, "reviewed_at": None,
        }

    def test_duplicate_order_within_group_is_rejected(self):
        rules = [
            self._rule("R1", "RF001", 1, "weight >= 80"),
            self._rule("R2", "RF001", 1, "breathlessness"),  # same order in same group
        ]
        errors, warnings = parser.validate_rule_groups(rules, FACTS, self._settings())
        assert len(errors) >= 1
        assert any("R1" in e or "R2" in e or "order" in e.lower() for e in errors)

    def test_duplicate_order_across_groups_is_accepted(self):
        """Order is group-scoped — same order value in different groups is valid."""
        rules = [
            self._rule("R1", "RF001", 1, "weight >= 80"),
            self._rule("R2", "RF002", 1, "breathlessness"),  # same order, DIFFERENT group
        ]
        # Need RF002 in red_flags too — the test doesn't call parse_rules so no validation
        errors, warnings = parser.validate_rule_groups(rules, FACTS, self._settings())
        assert errors == []

    def test_no_reading_rule_with_lower_order_than_clinical_rule_is_rejected(self):
        """A no_reading rule must have a HIGHER order number than all clinical rules."""
        rules = [
            self._rule("R1", "RF001", 1, "no_reading(weight, 7d)"),   # order 1 — too low
            self._rule("R2", "RF001", 2, "weight >= 80"),              # order 2 — clinical
        ]
        errors, warnings = parser.validate_rule_groups(rules, FACTS, self._settings())
        assert len(errors) >= 1
        assert any("no_reading" in e.lower() or "order" in e.lower() for e in errors)

    def test_two_no_reading_rules_both_above_clinical_is_accepted(self):
        """Multiple no_reading rules are allowed as long as all are above clinical rules."""
        rules = [
            self._rule("R1", "RF001", 1, "weight >= 80"),              # clinical
            self._rule("R2", "RF001", 2, "no_reading(weight, 7d)"),   # above clinical
            self._rule("R3", "RF001", 3, "no_reading(weight, 14d)"),  # also above clinical
        ]
        errors, warnings = parser.validate_rule_groups(rules, FACTS, self._settings())
        assert errors == []

    def test_cadence_longer_than_valid_for_warns(self):
        """min_days_between_checkins > shortest Required fact Valid For → warning."""
        # weight has valid_for_hours=168 (7d). Cadence "14d" (336h) > 168h → warn
        rules = [self._rule("R1", "RF001", 1, "weight >= 80")]
        errors, warnings = parser.validate_rule_groups(rules, FACTS, self._settings(cadence="14d"))
        assert errors == []   # warning, not error
        assert len(warnings) >= 1

    def test_count_threshold_unreachable_at_cadence_warns(self):
        """count(fact, window) >= N where N > window/cadence recordings → warning."""
        # count(breathlessness, 7d) >= 7 at 3d cadence: 7d/3d ≈ 2 recordings max → warn
        rules = [self._rule("R1", "RF001", 1, "count(breathlessness, 7d) >= 7")]
        errors, warnings = parser.validate_rule_groups(rules, FACTS, self._settings(cadence="3d"))
        assert errors == []
        assert len(warnings) >= 1

    def test_no_reading_straddling_clinical_rules_is_rejected(self):
        """A no_reading rule below clinical, even with another no_reading above, is rejected."""
        rules = [
            self._rule("R1", "RF001", 1, "no_reading(weight, 7d)"),   # order 1 — below clinical
            self._rule("R2", "RF001", 2, "weight >= 80"),              # order 2 — clinical
            self._rule("R3", "RF001", 3, "no_reading(weight, 14d)"),  # order 3 — above clinical
        ]
        errors, warnings = parser.validate_rule_groups(rules, FACTS, self._settings())
        assert len(errors) >= 1
